"""Build the six-page technical report (editable .docx). Numbers come from frontend/public/evaluation.json."""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SHOT, FIG = ROOT / "docs" / "screenshots", ROOT / "docs" / "figures"
EV = json.loads((ROOT / "frontend/public/evaluation.json").read_text(encoding="utf-8"))
NAVY, TEAL, GOLD, RED = RGBColor(0x12, 0x24, 0x3A), RGBColor(0x16, 0x7E, 0x80), RGBColor(0xC6, 0xA3, 0x6B), RGBColor(0xB9, 0x2D, 0x39)
BODY = float(sys.argv[1]) if len(sys.argv) > 1 else 9.5

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.7)
sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.4)
st = doc.styles["Normal"]
st.font.name, st.font.size = "Calibri", Pt(BODY)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after, st.paragraph_format.space_before, st.paragraph_format.line_spacing = Pt(3), Pt(0), 1.04


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_); tcPr.append(shd)


def para(text="", bold=False, size=None, color=None, italic=False, align=None, after=None, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        add_runs(p, text, bold, size, color, italic)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p


def add_runs(p, text, bold=False, size=None, color=None, italic=False):
    for i, part in enumerate(text.split("**")):  # **bold** inline
        if not part: continue
        r = p.add_run(part); r.bold = bold or (i % 2 == 1); r.italic = italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color


def h1(text, num):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{num}  "); r.font.size = Pt(18); r.font.color.rgb = GOLD; r.bold = True; r.font.name = "Cambria"
    r = p.add_run(text); r.font.size = Pt(18); r.font.color.rgb = NAVY; r.bold = True; r.font.name = "Cambria"
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "10"), ("space", "2"), ("color", "167E80")): bt.set(qn(f"w:{k}"), v)
    b.append(bt); pPr.append(b)


def h2(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(BODY + 2); r.font.color.rgb = TEAL; r.font.name = "Cambria"


def bullets(items, size=None):
    for it in items:
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(1); p.paragraph_format.left_indent = Cm(0.55)
        add_runs(p, it, size=size)


def table(rows, widths, head=True, size=None, zebra=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            add_runs(p, str(val), bold=(head and i == 0), size=size or BODY - 1, color=RGBColor(255, 255, 255) if head and i == 0 else None)
            if head and i == 0: shade(c, "12243A")
            elif zebra and i % 2 == 0: shade(c, "F7F3EB")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def pic(path, width, caption=None):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(path), width=Cm(width))
    if caption: para(caption, italic=True, size=BODY - 2, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=3)


def pic_row(items, width, cap_size=None):
    t = doc.add_table(rows=2, cols=len(items)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (path, cap) in enumerate(items):
        c = t.cell(0, j); c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraphs[0].add_run().add_picture(str(path), width=Cm(width))
        c2 = t.cell(1, j); c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(c2.paragraphs[0], cap, italic=True, size=BODY - 2, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# header / footer
hp = sec.header.paragraphs[0]; hp.text = ""
r = hp.add_run("WASL Arena  ·  Technical Report  ·  Team 204  ·  Code AYUCK  ·  Cycle 3 (grades 9-12)  ·  AI for Family"); r.font.size = Pt(8); r.font.color.rgb = TEAL
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("Page "); fr.font.size = Pt(8)
for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    run = fp.add_run(); run.font.size = Pt(8)
    if kind:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); run._r.append(e)
    else:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; run._r.append(e)
fr2 = fp.add_run(" of 6"); fr2.font.size = Pt(8)

pt, fx = EV["pytest"], EV["fixtures"]
gem = EV["transcription_check"]

# ------------------------------------------------------------------ PAGE 1
t = doc.add_table(rows=1, cols=1); c = t.cell(0, 0); shade(c, "12243A"); c.text = ""
p = c.paragraphs[0]; p.paragraph_format.space_before = Pt(6)
r = p.add_run("WASL Arena"); r.font.size = Pt(30); r.bold = True; r.font.color.rgb = RGBColor(0xF7, 0xF3, 0xEB); r.font.name = "Cambria"
p2 = c.add_paragraph(); r = p2.add_run("One family. Three games. One fair AI referee."); r.font.size = Pt(13); r.font.color.rgb = GOLD; r.font.name = "Cambria"
p3 = c.add_paragraph(); p3.paragraph_format.space_after = Pt(6)
r = p3.add_run("Family AI Game Day · Technical report · 6th Annual COE STEM Programming with AI Competition, Abu Dhabi University"); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0xDD, 0xE6, 0xEA)
p4 = c.add_paragraph(); p4.paragraph_format.space_after = Pt(6)
r = p4.add_run("Team 204  ·  Team code AYUCK  ·  Track: Cycle 3 (grades 9-12)"); r.font.size = Pt(10.5); r.bold = True; r.font.color.rgb = GOLD
p5 = c.add_paragraph(); p5.paragraph_format.space_after = Pt(6)
r = p5.add_run("May Ahmed Mourad (G12)  ·  Shamsa Faris Almazrouei (G12)  ·  Fatima Mohamed Almazrouei (G11)"); r.font.size = Pt(10); r.font.color.rgb = RGBColor(0xF7, 0xF3, 0xEB)
doc.add_paragraph().paragraph_format.space_after = Pt(2)

h1("Executive summary, problem and objectives", "1")
para("**WASL Arena** is a working web application that lets a family play three games, **football plus two original games**, while an AI referee watches the recorded video and listens to the audio, finds what actually happened, applies the written rules and keeps score. A human host runs the session: registering players and teams, running rounds and timers, and reviewing any call the AI is unsure about. The defining idea is **AI Referee Replay**: every score change is linked to evidence you can replay: the source video, timestamps, annotated frames, the rule applied, the points awarded and a verification status. If the AI cannot see it, it does not score it: the event is marked **NEEDS_REVIEW** and waits for the host.")
h2("The problem")
para("Family games are decided by arguments. “Was that a goal?” “Did you say it before you showed the card?” An AI judge could settle these fairly, but only if families can trust it. Black-box scoring fails that test, and so does an AI that guesses. The challenge asks for an app whose AI **observes real play through video and audio, applies the rules and keeps score**, with a human host in charge and a demo that shows the AI judging working on real gameplay.")
h2("Objectives")
bullets(["Deliver **exactly three games**: Football (required), **Falaj Flow** (team) and **Souq Signal** (individual), covering team and individual play and avoiding the example games excluded by the brief.",
         "Judge **real video/audio** from an upload or direct camera/microphone recording using computer vision and speech transcription, not simulated inputs.",
         "Make every judgment **explainable, verifiable and tamper-evident**, with uncertain detections routed to the host.",
         "Support the host journey end to end: choose a game, players and teams, rules and readiness, rounds and turns, AI judging, results with reasons and MVP, play another game."])
h2("How the brief is met")
table([["Official requirement", "WASL Arena implementation", "Status"],
       ["Football + 2 original games (Cycle 3)", "Football AI · Falaj Flow · Souq Signal", "Built"],
       ["Team and individual competition", "Football, Falaj = teams · Souq Signal = individual", "Built"],
       ["Upload video with audio or record from camera/mic", "File upload and MediaRecorder capture in the Live Arena", "Built"],
       ["AI follows players/ball; goals, possession, MVP", "Colour tracking, goal zones, observed-possession MVP", "Built; synthetic-tested"],
       ["Host sets up, runs rounds/turns/timers, shows reasons", "Setup wizard, turn order, timer, Replay with reasons", "Built"],
       ["AI report verified by code or hash", "SHA-256 + HMAC seal; stand-alone Python judge; import rejects edits", "Built; tested"],
       ["Demo with real gameplay", "Pipeline proven on synthetic fixtures; **real clips pending recording**", "Pending real footage"]],
      [5.2, 9.0, 3.4], size=BODY - 1.2)
page_break()

# ------------------------------------------------------------------ PAGE 2
h1("System concept, the three games and the user journey", "2")
para("WASL (وصل) means “connection”: the app connects physical family play to digital judging. The same flow serves all games: **1 Choose a game → 2 Players and teams → 3 Rules and ready → 4 Round in play → 5 AI judge results → 6 Final results**, exactly the host journey in the brief.")
table([["Game", "Mode", "What families do", "What the AI does", "Scoring"],
       ["**Football AI**", "Team", "Mini-match with a bright ball, two goals, contrasting shirts, fixed camera", "Tracks ball and players, detects goals, measures possession time per team and player", "Goal +1; tie → golden goal; MVP = longest possession"],
       ["**Falaj Flow** (فلج)", "Team", "Build a water-channel network on a printed 5×5 board with printed tiles in 90 s (inspired by the UAE falaj irrigation system)", "Reads every tile and rotation, rebuilds the network as a graph, traces water to three farms", "Farm +3; leak −1; speed bonus ≤ +2; tie → shortest network, time"],
       ["**Souq Signal** (إشارة السوق)", "Individual", "Shopping mission with a budget: pick product cards, show them and name them aloud (inspired by UAE souqs)", "Recognises cards, transcribes Arabic/English speech, checks audio and video agree, does the budget maths", "Item +2; spoken match +1; over budget −2; duplicates 0"]],
      [2.6, 1.6, 4.6, 4.6, 4.2], size=BODY - 1.4)
para("Each game exercises a different AI capability (motion tracking; visual recognition and graph reasoning; multimodal vision plus speech), which keeps the collection varied and gives judges three distinct demonstrations of the same fairness principle.")
h2("User journey (screens of the working prototype)")
pic_row([(SHOT / "02-setup.png", "Family setup: players, teams, language, match rules"), (SHOT / "03-library.png", "Game library: exactly three games")], 8.2)
pic_row([(SHOT / "04-prepare.png", "Preparation: rules, calibration, equipment, printables"), (SHOT / "05-arena.png", "Live Arena: camera, timer, turns, verified scoreboard")], 8.2)
para("Eight screens are implemented: **Launch, Family Setup, Game Library, Game Preparation, Live Arena, AI Referee Replay, Family Cup, Evaluation Lab**. The interface is bilingual (English/Arabic with right-to-left layout), and printable boards, tiles and product cards (with the exact markers the AI reads) are downloadable from the preparation screen.", size=BODY - 0.5)
page_break()

# ------------------------------------------------------------------ PAGE 3
h1("System architecture, technology and AI pipeline", "3")
pic(FIG / "architecture.png", 17.2)
table([["Layer", "Technology", "Role"],
       ["Frontend", "React 19, TypeScript, Vite, Tailwind CSS 4, React Router, Three.js", "Eight responsive screens; MediaRecorder capture; Canvas overlays synchronised with video; original reactive 3D AI-Referee orb with CSS fallback"],
       ["Backend", "Python, FastAPI, Pydantic, Uvicorn", "Upload, background judging jobs with real stage reporting, report store, host review, audit chain"],
       ["Vision / media", "OpenCV (HSV segmentation, contours, ArUco, homography), NumPy, FFmpeg", "Frame sampling at 4–10 fps, object tracking, marker decoding, audio extraction"],
       ["Speech", "Gemini API (server-side key)", "**Transcription only**: timed text segments; an energy gate drops segments with no real sound"],
       ["Rules", "Pure-Python, versioned rule engines", "WASL-FOOTBALL-R1.0, WASL-FALAJ-R1.0, WASL-SOUQ-R1.0: the only scoring authority"],
       ["Delivery", "GitHub → Vercel (frontend); Render (backend)", "Backend URL via VITE_API_URL; credentials only in server environment variables"]],
      [2.3, 6.0, 9.3], size=BODY - 1.4)
pic(FIG / "pipeline.png", 17.2)
h2("Evidence integrity")
para("Each job is hashed at upload (**SHA-256 of the original file**). The finished report carries the analysed-file hash, the rule version, event timestamps, bounding boxes, annotated evidence frames and validation checks, then is **sealed with a SHA-256 digest and HMAC signature**. Anyone holding the key can detect a single changed character; the app rejects tampered reports on import. Host decisions recompute the score through the same rule engine, are appended to a **hash-chained audit log**, and re-seal the report while preserving the original AI seal. The stand-alone judge (`python -m app.judge_cli`) produces the same sealed report outside the app, satisfying the brief's second connection option.", size=BODY - 0.3)
page_break()

# ------------------------------------------------------------------ PAGE 4
h1("Football, Falaj Flow and Souq Signal: AI methodologies", "4")
h2("Football AI: tracking, goals, possession and MVP")
para("Frames are sampled at 10 fps and downscaled to 640 px. The **ball** is the best circular contour inside a calibrated HSV colour range (area and circularity gated, nearest to a velocity-predicted position). **Players** are team-colour blobs (morphological open/close, area filter) tracked by greedy nearest-neighbour association with 3-second re-association and a per-team player cap. A **goal** is the ball observed inside the calibrated goal polygon for ≥ 3 consecutive observed samples after being seen outside it, with a 3 s cooldown; a goal whose entry was not observed, or that lasts fewer samples, is NEEDS_REVIEW. A ball that heads for a goal and then disappears becomes a **POSSIBLE_GOAL** (never auto-scored). **Possession** credits the nearest player within 0.45 body-lengths for ≥ 3 samples; contested moments and unobserved time are credited to nobody. **MVP** is the player with the longest verified possession (as defined in the brief); because shirt numbers are not read, the host assigns colour tracks to registered players.")
pic_row([(ROOT / "frontend/public/samples/evidence/football_goal/goal-2-moment.jpg", "Football goal evidence frame (synthetic fixture)"),
         (ROOT / "frontend/public/samples/evidence/falaj_valid/final-board.jpg", "Falaj: board read from video, water path (teal)")], 8.4)
h2("Falaj Flow: board recognition, graph reconstruction and validation")
para("Four ArUco corner markers define a homography from the camera image to the 5×5 board. Each tile carries an ArUco marker whose **id gives the tile type** (straight, elbow, tee, cross) and whose **in-plane rotation gives the tile rotation** in 90° steps. A board arrangement counts only when it is **stable for 1 s**, so hands moving over tiles are ignored. The arrangement becomes a graph; breadth-first search from the well follows only open sides whose neighbour has a matching open side. An open end that meets the board edge, a rock, an empty cell or a closed side is a **leak**. Output: farms reached (+3 each), leaks (−1), time bonus by completion time, network size for the tie-break, and a digital twin that animates water only along rule-validated edges.")
h2("Souq Signal: multimodal verification")
para("Each product card carries a unique marker (two copies per product); sightings are merged into **presentation intervals** (≥ 0.8 s, gaps ≤ 2 s). Speech is transcribed server-side, then matched to bilingual product aliases (Arabic normalisation, ال/و prefixes) and timestamped inside the segment. The rule engine then checks: requested item shown (+2), spoken within ±5 s (+1), **audio ≠ video mismatch → host review**, duplicate card (0), and distinct card prices against the budget (−2 once). With no transcript, spoken-match points stay NEEDS_REVIEW. Gemini never decides points: it provides text; deterministic code scores.", size=BODY - 0.2)
page_break()

# ------------------------------------------------------------------ PAGE 5
h1("Testing, results, validation and limitations", "5")
para(f"All results below were produced by running the code (generated {EV['generated_at'][:10]}). **{pt['passed']} of {pt['total']} automated tests passed, {pt['failed']} failed.** Tests cover the three rule engines, the full CV-plus-rules pipelines, report sealing, tamper detection, the upload → job → report → host-review API flow and the audit chain.")
rows = [["Fixture (synthetic render)", "Ground truth", "Pipeline output", "Result"]]
for f in fx:
    rows.append([f["name"], ", ".join(f"{k}={v}" for k, v in f["expected"].items() if k != "status"), ", ".join(f"{k}={v}" for k, v in f["actual"].items()), "PASS" if f["pass"] else "FAIL"])
table(rows, [3.6, 5.2, 7.2, 1.6], size=BODY - 1.8)
by_file = {}
for c in pt["cases"]:
    by_file.setdefault(c["file"], [0, 0]); by_file[c["file"]][1] += 1; by_file[c["file"]][0] += c["status"] == "passed"
table([["Test module", "Covers", "Passed"]] + [[k.replace("test_", ""), {"football_rules": "valid goal, no goal, uncertain, score update, tie-break, MVP, penalties", "falaj_rules": "complete route, partial, blocked, invalid connection, scoring, tie-break", "souq_rules": "correct, incorrect, mismatch, budget, duplicate, Arabic matching, tie-break", "pipelines": "end-to-end video → events → score for all three games", "integrity_api": "seal/tamper, API flow, host review recompute, import rejection, audit chain"}.get(k.replace("test_", ""), ""), f"{a}/{b}"] for k, (a, b) in by_file.items()], [3.6, 11.4, 2.6], size=BODY - 1.6)
h2("Speech check and what the tests found")
if gem and gem.get("status") == "OK":
    para("A live call to the transcription model on synthetic text-to-speech audio returned: " + "; ".join(f"“{s['text']}” ({s['start']}–{s['end']} s)" for s in gem["segments"]) + f". Deterministic scoring of that transcript gave **{gem['total']}** points (expected {gem['expected_total']}). Testing exposed three real defects that were fixed: the well’s own inlet was mis-counted as a leak; Arabic words with the ال prefix were not matched; and a host-accepted “possible goal” was not scored.", size=BODY - 0.3)
h2("Limitations and honesty")
bullets(["**No real family footage has been analysed yet.** Fixtures are synthetic renders with known ground truth; they test the pipelines, not real-world accuracy. No accuracy percentage is claimed. Real clips are to be recorded and reported separately.",
         "Football assumes a fixed camera, single-colour ball, contrasting shirts and calibrated goals; colour segmentation is lighting-sensitive; shirt numbers are not read.",
         "Falaj/Souq read printed markers; a hidden or damaged marker counts as absent. Speech timing is approximate (≈ ±1 s).",
         "Detection scores are heuristic indicators, not calibrated probabilities."], size=BODY - 0.5)
page_break()

# ------------------------------------------------------------------ PAGE 6
h1("Fairness, privacy, deployment, impact and conclusions", "6")
h2("Fairness and transparency")
bullets(["**Rules are written, versioned and deterministic.** The AI proposes events; the rule engine decides points, so the same evidence always gives the same score.",
         "**Uncertainty is visible.** Unobserved or ambiguous moments become NEEDS_REVIEW and score nothing until the host accepts them; unobserved ball time is shown as “unknown possession”, not guessed.",
         "**Every call is contestable:** the host can accept, reject or clear any event; the score is recomputed and the decision is logged with name and time in a hash-chained audit trail.",
         "**Inclusive by design:** team and individual formats, Arabic/English interface and speech, heritage themes (falaj, souq) and equipment families can print at home."], size=BODY - 0.3)
h2("Privacy and security")
bullets(["Video is processed only to produce the judging report; uploads are limited in size and stored on the judge server, not shared with third parties other than the transcription call (audio only) when a key is configured.",
         "No faces are identified or stored as biometric data: players are colour tracks mapped to names the host types. Families should record only their own game and obtain consent from everyone filmed.",
         "All credentials stay server-side in environment variables (.env is git-ignored); the browser never sees them. Reports are signed to detect tampering."], size=BODY - 0.3)
h2("Deployment")
para("The repository **wasl-arena** holds frontend, backend, tests, tools, documentation and AI-usage records. The React app deploys from GitHub to Vercel (`frontend/vercel.json`, `VITE_API_URL`); the FastAPI judge deploys to Render (`render.yaml`) because video processing needs a long-running Python runtime with FFmpeg and OpenCV. Production build verified locally (TypeScript clean, Vite build). Live URLs are recorded in the README once the accounts are connected; publication is not claimed here until verified.", size=BODY - 0.3)
h2("Impact and future work")
para("WASL Arena shows that AI can make family gaming **more interactive** (real physical play), **more inclusive** (Arabic and English, team and solo), **culturally meaningful** (falaj and souq) and **transparently fair** (replayable evidence). Next steps: record and evaluate real family footage per game; add pretrained detectors and jersey-number reading to reduce calibration needs; add whistle/“goal!” audio cues to Football; and let hosts define their own games from the same rule-engine template.", size=BODY - 0.3)
h2("Conclusion")
para("A fair referee is not one that is always right: it is one that **shows its evidence, admits doubt and lets people decide**. WASL Arena delivers that for three games, with a working prototype, a verified pipeline and honest limits.", size=BODY - 0.3)
h2("References")
para("[1] Abu Dhabi University College of Engineering, “Family AI Game Day: Competition Questions – Cycle 3 Track (2026)”.  [2] OpenCV ArUco marker detection documentation.  [3] FFmpeg documentation.  [4] FastAPI, React, Vite and Three.js documentation.  [5] Google Gemini API, audio understanding.  [6] RFC 2104, HMAC: Keyed-Hashing for Message Authentication.  [7] UAE falaj irrigation systems, UNESCO World Heritage listing (Al Ain cultural sites).", size=BODY - 1.5)

out = ROOT / "docs" / "WASL-Arena-Technical-Report.docx"
doc.save(out)
print("saved", out)
